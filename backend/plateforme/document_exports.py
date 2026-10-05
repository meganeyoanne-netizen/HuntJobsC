import re
import unicodedata
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor, white
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    FrameBreak,
    HRFlowable,
    KeepTogether,
    PageTemplate,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


THEMES = {
    "cv-classique": {"accent": "1E3A8A", "font": "Arial", "pdf_font": "Helvetica", "title_align": "center"},
    "cv-moderne": {"accent": "123C57", "font": "Arial", "pdf_font": "Helvetica", "title_align": "left"},
    "cv-creatif": {"accent": "312E81", "font": "Arial", "pdf_font": "Helvetica", "title_align": "left"},
    "portfolio-minimal": {"accent": "0F172A", "font": "Arial", "pdf_font": "Helvetica", "title_align": "left"},
    "portfolio-projets": {"accent": "0F2345", "font": "Arial", "pdf_font": "Helvetica", "title_align": "left"},
    "portfolio-editorial": {"accent": "78350F", "font": "Georgia", "pdf_font": "Times-Roman", "title_align": "left"},
    "lettre-classique": {"accent": "1F2937", "font": "Georgia", "pdf_font": "Times-Roman", "title_align": "left"},
    "lettre-moderne": {"accent": "1D4ED8", "font": "Arial", "pdf_font": "Helvetica", "title_align": "left"},
    "lettre-elegante": {"accent": "6B2148", "font": "Georgia", "pdf_font": "Times-Roman", "title_align": "center"},
    "libre": {"accent": "4F46E5", "font": "Arial", "pdf_font": "Helvetica", "title_align": "left"},
}


def _theme(design):
    return THEMES.get(design, THEMES["libre"])


def _normalise(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


def _text(value, keys=()):
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, dict):
        for key in keys:
            if value.get(key):
                return str(value[key]).strip()
    return ""


def _blocks(content):
    blocks = []
    for raw in str(content or "").replace("```markdown", "").replace("```", "").splitlines():
        line = raw.strip()
        if not line:
            blocks.append(("space", ""))
        elif line.startswith("### "):
            blocks.append(("subheading", line[4:].strip()))
        elif line.startswith("## "):
            blocks.append(("heading", line[3:].strip()))
        elif line.startswith("# "):
            blocks.append(("title", line[2:].strip()))
        elif re.match(r"^[-*•]\s+", line):
            blocks.append(("bullet", re.sub(r"^[-*•]\s+", "", line)))
        elif re.match(r"^\*\*.+\*\*:?$", line):
            blocks.append(("heading", line.strip("*: ")))
        else:
            blocks.append(("body", line))
    compact = []
    for block in blocks:
        if block[0] != "space" or (compact and compact[-1][0] != "space"):
            compact.append(block)
    return compact


def _markdown_sections(content):
    sections = {"intro": []}
    current = "intro"
    aliases = {
        "profil": "profile", "resume": "profile", "a propos": "profile", "objectif": "profile",
        "experience": "experience", "experiences": "experience", "parcours professionnel": "experience",
        "formation": "education", "formations": "education", "education": "education",
        "competences": "skills", "points forts": "skills", "mes points forts": "skills",
        "langues": "languages", "interets": "interests", "centres d interet": "interests",
    }
    for kind, value in _blocks(content):
        if kind in {"heading", "title"}:
            key = _normalise(value)
            current = next((target for label, target in aliases.items() if label in key), key or "other")
            sections.setdefault(current, [])
        elif kind != "space":
            sections.setdefault(current, []).append((kind, value))
    return sections


def _section_text(sections, key):
    return " ".join(value for kind, value in sections.get(key, []) if kind in {"body", "bullet"}).strip()


def _profile_snapshot(generation):
    source = generation.source_data or {}
    profile = source.get("profil") or {}
    if not profile and hasattr(generation.user, "profil_candidat"):
        candidate = generation.user.profil_candidat
        profile = {
            "titre_professionnel": candidate.titre_professionnel,
            "bio": candidate.bio,
            "localisation": candidate.localisation,
            "competences": candidate.competences,
            "experiences": candidate.experiences,
            "formations": candidate.formations,
            "langues": candidate.langues,
            "linkedin_url": candidate.linkedin_url,
            "github_url": candidate.github_url,
            "portfolio_url": candidate.portfolio_url,
        }
    user_data = profile.get("user") if isinstance(profile.get("user"), dict) else {}
    return source, profile, user_data


def _cv_data(generation):
    source, profile, user_data = _profile_snapshot(generation)
    sections = _markdown_sections(generation.contenu)
    name = source.get("nom") or generation.user.full_name or generation.titre
    email = source.get("email") or user_data.get("email") or generation.user.email
    phone = user_data.get("telephone") or getattr(generation.user, "telephone", "")
    skills = profile.get("competences") or [value for _, value in sections.get("skills", [])]
    languages = profile.get("langues") or [value for _, value in sections.get("languages", [])]
    experiences = profile.get("experiences") or [{"position": value} for _, value in sections.get("experience", [])]
    education = profile.get("formations") or [{"degree": value} for _, value in sections.get("education", [])]
    return {
        "name": name,
        "initials": "".join(part[0].upper() for part in name.split()[:2] if part),
        "title": profile.get("titre_professionnel") or "Profil professionnel",
        "email": email,
        "phone": phone,
        "location": profile.get("localisation", ""),
        "bio": _section_text(sections, "profile") or profile.get("bio") or _section_text(sections, "intro"),
        "skills": [str(item) for item in skills if str(item).strip()],
        "languages": [_text(item, ("name", "langue", "language")) or str(item) for item in languages if str(item).strip()],
        "experiences": experiences,
        "education": education,
        "links": [profile.get(key) for key in ("linkedin_url", "github_url", "portfolio_url") if profile.get(key)],
    }


def _experience(item):
    if not isinstance(item, dict):
        return str(item), "", "", ""
    role = _text(item, ("position", "poste", "role", "title"))
    company = _text(item, ("company", "entreprise", "organisation"))
    start = _text(item, ("startDate", "date_debut", "start"))
    end = "Aujourd'hui" if item.get("current") else _text(item, ("endDate", "date_fin", "end"))
    period = " - ".join(value for value in (start, end) if value)
    description = _text(item, ("description", "details", "missions"))
    return role or company or "Expérience professionnelle", company, period, description


def _education(item):
    if not isinstance(item, dict):
        return str(item), "", ""
    degree = _text(item, ("degree", "diplome", "formation", "title"))
    school = _text(item, ("school", "ecole", "etablissement", "institution"))
    start = _text(item, ("startDate", "date_debut", "start"))
    end = _text(item, ("endDate", "date_fin", "end"))
    return degree or school or "Formation", school, " - ".join(value for value in (start, end) if value)


def _shade_cell(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_margins(cell, top=120, start=160, bottom=120, end=160):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_border(cell, **edges):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge, config in edges.items():
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        for key, value in config.items():
            node.set(qn(f"w:{key}"), str(value))


def _remove_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(cell, top={"val": "nil"}, left={"val": "nil"}, bottom={"val": "nil"}, right={"val": "nil"}, insideH={"val": "nil"}, insideV={"val": "nil"})


def _doc_run(paragraph, text, size=10, color="1F2937", bold=False, font="Arial"):
    run = paragraph.add_run(str(text or ""))
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    return run


def _doc_heading(container, text, color, font="Arial", size=9, centered=False):
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    _doc_run(paragraph, str(text).upper(), size=size, color=color, bold=True, font=font)
    return paragraph


def _doc_body(container, text, color="334155", font="Arial", size=9.5, before=0, after=4):
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.05
    _doc_run(paragraph, text, size=size, color=color, font=font)
    return paragraph


def _doc_contact_line(data):
    return "  |  ".join(value for value in (data["email"], data["phone"], data["location"]) if value)


def _add_doc_experiences(container, data, accent, font="Arial", compact=False):
    for item in data["experiences"][:5]:
        role, company, period, description = _experience(item)
        paragraph = container.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.keep_with_next = True
        _doc_run(paragraph, role, size=9.5 if compact else 10, color="1F2937", bold=True, font=font)
        if period:
            _doc_run(paragraph, "  " + period, size=8, color="94A3B8", font=font)
        if company:
            _doc_body(container, company, color=accent, font=font, size=8.5, after=1)
        if description:
            _doc_body(container, description, font=font, size=8.5 if compact else 9, after=3)


def _add_doc_education(container, data, accent, font="Arial"):
    for item in data["education"][:4]:
        degree, school, period = _education(item)
        paragraph = container.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(1)
        _doc_run(paragraph, degree, size=9.5, color="1F2937", bold=True, font=font)
        if period:
            _doc_run(paragraph, "  " + period, size=8, color="94A3B8", font=font)
        if school:
            _doc_body(container, school, color=accent, font=font, size=8.5, after=2)


def _base_docx(theme):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.55)
    section.left_margin = section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]
    normal.font.name = theme["font"]
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    normal.paragraph_format.space_after = Pt(3)
    return document


def _docx_cv_classique(generation):
    data = _cv_data(generation)
    document = _base_docx(_theme("cv-classique"))
    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2)
    _doc_run(name, data["name"], size=20, color="172554", bold=True)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    _doc_run(title, data["title"].upper(), size=8, color="64748B")
    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(8)
    _doc_run(contact, _doc_contact_line(data), size=8, color="64748B")
    ppr = contact._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:color"), "1E3A8A")
    borders.append(bottom)
    ppr.append(borders)
    if data["bio"]:
        _doc_heading(document, "Profil", "1E3A8A")
        _doc_body(document, data["bio"])
    if data["experiences"]:
        _doc_heading(document, "Expérience", "1E3A8A")
        _add_doc_experiences(document, data, "1E3A8A")
    if data["education"]:
        _doc_heading(document, "Formation", "1E3A8A")
        _add_doc_education(document, data, "1E3A8A")
    if data["skills"]:
        _doc_heading(document, "Compétences", "1E3A8A")
        _doc_body(document, "  |  ".join(data["skills"][:12]))
    if data["languages"]:
        _doc_heading(document, "Langues", "1E3A8A")
        _doc_body(document, "  |  ".join(data["languages"][:8]))
    return document


def _docx_cv_moderne(generation):
    data = _cv_data(generation)
    document = _base_docx(_theme("cv-moderne"))
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.35)
    section.left_margin = section.right_margin = Inches(0.35)
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.05)
    table.columns[1].width = Inches(5.55)
    _remove_table_borders(table)
    side, main = table.rows[0].cells
    side.width, main.width = Inches(2.05), Inches(5.55)
    side.vertical_alignment = main.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _set_cell_margins(side, 300, 220, 260, 220)
    _set_cell_margins(main, 300, 300, 260, 260)
    _shade_cell(side, "123C57")
    avatar = side.paragraphs[0]
    avatar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    avatar.paragraph_format.space_after = Pt(12)
    _doc_run(avatar, data["initials"] or "CV", size=18, color="FFFFFF", bold=True)
    _doc_heading(side, "Contact", "DFF5FA", size=8)
    for value in (data["email"], data["phone"], data["location"]):
        if value:
            _doc_body(side, value, color="DFF5FA", size=8, after=3)
    if data["skills"]:
        _doc_heading(side, "Compétences", "DFF5FA", size=8)
        for skill in data["skills"][:10]:
            _doc_body(side, skill, color="DFF5FA", size=8, after=2)
    if data["languages"]:
        _doc_heading(side, "Langues", "DFF5FA", size=8)
        for language in data["languages"][:6]:
            _doc_body(side, language, color="DFF5FA", size=8, after=2)
    name = main.paragraphs[0]
    name.paragraph_format.space_after = Pt(3)
    _doc_run(name, data["name"], size=22, color="123C57", bold=True)
    _doc_body(main, data["title"].upper(), color="256B83", size=8.5, after=10)
    if data["bio"]:
        _doc_heading(main, "Profil", "123C57")
        _doc_body(main, data["bio"], size=9)
    if data["experiences"]:
        _doc_heading(main, "Expérience", "123C57")
        _add_doc_experiences(main, data, "2D9CAA", compact=True)
    if data["education"]:
        _doc_heading(main, "Formation", "123C57")
        _add_doc_education(main, data, "2D9CAA")
    return document


def _docx_cv_creatif(generation):
    data = _cv_data(generation)
    document = _base_docx(_theme("cv-creatif"))
    header = document.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width, header.columns[1].width = Inches(5.6), Inches(1.4)
    _remove_table_borders(header)
    left, badge = header.rows[0].cells
    _shade_cell(left, "312E81")
    _shade_cell(badge, "312E81")
    _set_cell_margins(left, 280, 300, 260, 220)
    _set_cell_margins(badge, 350, 80, 300, 80)
    intro = left.paragraphs[0]
    _doc_run(intro, "BONJOUR, JE SUIS", size=8, color="DDD6FE")
    name = left.add_paragraph()
    name.paragraph_format.space_after = Pt(2)
    _doc_run(name, data["name"], size=22, color="FFFFFF", bold=True)
    _doc_body(left, data["title"].upper(), color="DDD6FE", size=8.5, after=0)
    mark = badge.paragraphs[0]
    mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _doc_run(mark, "●", size=38, color="A78BFA")
    if data["bio"]:
        _doc_heading(document, "Mon parcours", "7C3AED")
        _doc_body(document, data["bio"], color="4C1D95")
    if data["experiences"]:
        card = document.add_table(rows=1, cols=1)
        _remove_table_borders(card)
        cell = card.cell(0, 0)
        _shade_cell(cell, "F5F3FF")
        _set_cell_margins(cell, 180, 220, 180, 220)
        _doc_heading(cell, "Expérience récente", "4C1D95", size=8)
        _add_doc_experiences(cell, {**data, "experiences": data["experiences"][:3]}, "7C3AED", compact=True)
    if data["skills"]:
        _doc_heading(document, "Mes points forts", "7C3AED")
        columns = min(3, len(data["skills"]))
        chips = document.add_table(rows=(len(data["skills"][:12]) + columns - 1) // columns, cols=columns)
        _remove_table_borders(chips)
        for index, skill in enumerate(data["skills"][:12]):
            cell = chips.cell(index // columns, index % columns)
            _shade_cell(cell, "DDD6FE")
            _set_cell_margins(cell, 90, 120, 90, 120)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _doc_run(cell.paragraphs[0], skill, size=8, color="4C1D95", bold=True)
    if data["education"]:
        _doc_heading(document, "Formation", "7C3AED")
        _add_doc_education(document, data, "7C3AED")
    return document


def _docx_cv_libre(generation):
    data = _cv_data(generation)
    document = _base_docx(_theme("libre"))
    header = document.add_table(rows=1, cols=1)
    _remove_table_borders(header)
    cell = header.cell(0, 0)
    _shade_cell(cell, "DBEAFE")
    _set_cell_margins(cell, 220, 260, 200, 260)
    name = cell.paragraphs[0]
    _doc_run(name, data["name"], size=20, color="1D4ED8", bold=True)
    _doc_body(cell, data["title"].upper(), color="3B82F6", size=8.5, after=1)
    _doc_body(cell, _doc_contact_line(data), color="64748B", size=8, after=0)
    summary = document.add_table(rows=1, cols=2)
    _remove_table_borders(summary)
    left, right = summary.rows[0].cells
    for target, color in ((left, "E9D5FF"), (right, "BFDBFE")):
        _shade_cell(target, color)
        _set_cell_margins(target, 180, 180, 180, 180)
    _doc_heading(left, "Profil", "6D28D9", size=8)
    _doc_body(left, data["bio"] or "Présentation professionnelle", color="4C1D95", size=8.5)
    _doc_heading(right, "Compétences", "1D4ED8", size=8)
    _doc_body(right, " | ".join(data["skills"][:8]) or "Compétences du profil", color="1E3A8A", size=8.5)
    if data["experiences"]:
        card = document.add_table(rows=1, cols=1)
        _remove_table_borders(card)
        experience_cell = card.cell(0, 0)
        _shade_cell(experience_cell, "DCFCE7")
        _set_cell_margins(experience_cell, 180, 220, 180, 220)
        _doc_heading(experience_cell, "Expérience", "15803D", size=8)
        _add_doc_experiences(experience_cell, data, "16A34A", compact=True)
    if data["education"]:
        _doc_heading(document, "Formation", "4F46E5")
        _add_doc_education(document, data, "4F46E5")
    return document


def _docx_generic(generation):
    theme = _theme(generation.design)
    document = _base_docx(theme)
    accent = theme["accent"]
    header = document.add_table(rows=1, cols=1)
    _remove_table_borders(header)
    cell = header.cell(0, 0)
    _shade_cell(cell, accent)
    _set_cell_margins(cell, 240, 260, 220, 260)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if theme["title_align"] == "center" else WD_ALIGN_PARAGRAPH.LEFT
    _doc_run(paragraph, generation.user.full_name or generation.titre, size=21, color="FFFFFF", bold=True, font=theme["font"])
    if generation.user.email:
        sub = cell.add_paragraph()
        sub.alignment = paragraph.alignment
        _doc_run(sub, generation.user.email, size=8.5, color="E6EEF8", font=theme["font"])
    for kind, text in _blocks(generation.contenu):
        if kind == "space":
            continue
        if kind == "title" and _normalise(text) in _normalise(generation.user.full_name or generation.titre):
            continue
        if kind in {"title", "heading", "subheading"}:
            _doc_heading(document, text, accent, font=theme["font"], size=11 if kind == "title" else 9)
        elif kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            _doc_run(paragraph, text, size=9.5, font=theme["font"])
        else:
            _doc_body(document, text, font=theme["font"])
    return document


def build_docx(generation):
    builders = {
        "cv-classique": _docx_cv_classique,
        "cv-moderne": _docx_cv_moderne,
        "cv-creatif": _docx_cv_creatif,
        "libre": _docx_cv_libre if generation.outil == "generer-cv" else _docx_generic,
    }
    document = builders.get(generation.design, _docx_generic)(generation)
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _xml(text):
    return str(text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _pdf_styles(accent, font="Helvetica"):
    styles = getSampleStyleSheet()
    bold = "Times-Bold" if font.startswith("Times") else "Helvetica-Bold"
    return {
        "name": ParagraphStyle("CVName", parent=styles["Normal"], fontName=bold, fontSize=20, leading=22, textColor=HexColor(accent), spaceAfter=3),
        "title": ParagraphStyle("CVTitle", parent=styles["Normal"], fontName=font, fontSize=8.5, leading=11, textColor=HexColor("#64748B"), spaceAfter=4),
        "heading": ParagraphStyle("CVHeading", parent=styles["Normal"], fontName=bold, fontSize=9, leading=11, textColor=HexColor(accent), spaceBefore=7, spaceAfter=4, keepWithNext=True),
        "body": ParagraphStyle("CVBody", parent=styles["Normal"], fontName=font, fontSize=9, leading=11.5, textColor=HexColor("#334155"), spaceAfter=4),
        "small": ParagraphStyle("CVSmall", parent=styles["Normal"], fontName=font, fontSize=7.8, leading=10, textColor=HexColor("#64748B"), spaceAfter=3),
        "white_heading": ParagraphStyle("CVWhiteHeading", parent=styles["Normal"], fontName=bold, fontSize=8, leading=10, textColor=white, spaceBefore=7, spaceAfter=4),
        "white_body": ParagraphStyle("CVWhiteBody", parent=styles["Normal"], fontName=font, fontSize=7.8, leading=10, textColor=HexColor("#DFF5FA"), spaceAfter=3),
    }


def _pdf_experiences(data, styles, accent, limit=5):
    flowables = []
    for item in data["experiences"][:limit]:
        role, company, period, description = _experience(item)
        heading = f"<b>{_xml(role)}</b>"
        if period:
            heading += f" <font color='#94A3B8' size='7.5'>{_xml(period)}</font>"
        block = [Paragraph(heading, styles["body"])]
        if company:
            block.append(Paragraph(f"<font color='{accent}'><b>{_xml(company)}</b></font>", styles["small"]))
        if description:
            block.append(Paragraph(_xml(description), styles["small"]))
        flowables.extend(block)
    return flowables


def _pdf_education(data, styles, accent):
    flowables = []
    for item in data["education"][:4]:
        degree, school, period = _education(item)
        text = f"<b>{_xml(degree)}</b>"
        if period:
            text += f" <font color='#94A3B8' size='7.5'>{_xml(period)}</font>"
        if school:
            text += f"<br/><font color='{accent}'>{_xml(school)}</font>"
        flowables.append(Paragraph(text, styles["body"]))
    return flowables


def _pdf_cv_classique(generation):
    data = _cv_data(generation)
    styles = _pdf_styles("#1E3A8A")
    styles["name"].alignment = TA_CENTER
    styles["title"].alignment = TA_CENTER
    contact_style = ParagraphStyle("CVContact", parent=styles["small"], alignment=TA_CENTER)
    story = [Paragraph(_xml(data["name"]), styles["name"]), Paragraph(_xml(data["title"].upper()), styles["title"]), Paragraph(_xml(_doc_contact_line(data)), contact_style), HRFlowable(width="100%", thickness=1.2, color=HexColor("#1E3A8A"), spaceBefore=2, spaceAfter=5)]
    if data["bio"]:
        story += [Paragraph("PROFIL", styles["heading"]), Paragraph(_xml(data["bio"]), styles["body"])]
    if data["experiences"]:
        story.append(Paragraph("EXPÉRIENCE", styles["heading"]))
        story.extend(_pdf_experiences(data, styles, "#1E3A8A"))
    if data["education"]:
        story.append(Paragraph("FORMATION", styles["heading"]))
        story.extend(_pdf_education(data, styles, "#1E3A8A"))
    if data["skills"]:
        story += [Paragraph("COMPÉTENCES", styles["heading"]), Paragraph(_xml("  |  ".join(data["skills"][:12])), styles["body"])]
    if data["languages"]:
        story += [Paragraph("LANGUES", styles["heading"]), Paragraph(_xml("  |  ".join(data["languages"][:8])), styles["body"])]
    buffer = BytesIO()
    pdf = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=.65*inch, leftMargin=.65*inch, topMargin=.5*inch, bottomMargin=.5*inch, title=generation.titre, author=data["name"])
    pdf.build(story)
    buffer.seek(0)
    return buffer


def _pdf_cv_moderne(generation):
    data = _cv_data(generation)
    styles = _pdf_styles("#123C57")
    buffer = BytesIO()
    document = BaseDocTemplate(buffer, pagesize=letter, leftMargin=0, rightMargin=0, topMargin=0, bottomMargin=0, title=generation.titre, author=data["name"])
    left_frame = Frame(.35*inch, .35*inch, 1.85*inch, 10.3*inch, leftPadding=10, rightPadding=10, topPadding=12, bottomPadding=12, id="sidebar")
    right_frame = Frame(2.4*inch, .45*inch, 5.65*inch, 10.1*inch, leftPadding=8, rightPadding=8, topPadding=10, bottomPadding=10, id="main")
    def background(canvas, _doc):
        canvas.saveState()
        canvas.setFillColor(HexColor("#123C57"))
        canvas.rect(.25*inch, .25*inch, 2.05*inch, 10.5*inch, fill=1, stroke=0)
        canvas.restoreState()
    document.addPageTemplates(PageTemplate(id="modern", frames=[left_frame, right_frame], onPage=background))
    side = [Paragraph(_xml(data["initials"] or "CV"), ParagraphStyle("Avatar", parent=styles["name"], alignment=TA_CENTER, fontSize=18, textColor=white, spaceAfter=13)), Paragraph("CONTACT", styles["white_heading"])]
    for value in (data["email"], data["phone"], data["location"]):
        if value:
            side.append(Paragraph(_xml(value), styles["white_body"]))
    if data["skills"]:
        side.append(Paragraph("COMPÉTENCES", styles["white_heading"]))
        side.extend(Paragraph(_xml(skill), styles["white_body"]) for skill in data["skills"][:10])
    if data["languages"]:
        side.append(Paragraph("LANGUES", styles["white_heading"]))
        side.extend(Paragraph(_xml(language), styles["white_body"]) for language in data["languages"][:6])
    main = [FrameBreak(), Paragraph(_xml(data["name"]), styles["name"]), Paragraph(_xml(data["title"].upper()), styles["title"])]
    if data["bio"]:
        main += [Paragraph("PROFIL", styles["heading"]), Paragraph(_xml(data["bio"]), styles["body"])]
    if data["experiences"]:
        main.append(Paragraph("EXPÉRIENCE", styles["heading"]))
        main.extend(_pdf_experiences(data, styles, "#2D9CAA", limit=4))
    if data["education"]:
        main.append(Paragraph("FORMATION", styles["heading"]))
        main.extend(_pdf_education(data, styles, "#2D9CAA"))
    document.build(side + main)
    buffer.seek(0)
    return buffer


def _pdf_cv_creatif(generation):
    data = _cv_data(generation)
    styles = _pdf_styles("#7C3AED")
    buffer = BytesIO()
    document = BaseDocTemplate(buffer, pagesize=letter, title=generation.titre, author=data["name"])
    frame = Frame(.65*inch, .5*inch, 7.2*inch, 7.75*inch, leftPadding=0, rightPadding=0, topPadding=6, bottomPadding=6, id="creative")
    def background(canvas, _doc):
        canvas.saveState()
        canvas.setFillColor(HexColor("#312E81"))
        canvas.rect(.55*inch, 8.55*inch, 7.4*inch, 1.9*inch, fill=1, stroke=0)
        canvas.setFillColor(HexColor("#A78BFA"))
        canvas.circle(7.05*inch, 9.45*inch, .38*inch, fill=1, stroke=0)
        canvas.setFillColor(HexColor("#DDD6FE"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(.82*inch, 10.0*inch, "BONJOUR, JE SUIS")
        canvas.setFillColor(white)
        canvas.setFont("Helvetica-Bold", 22)
        canvas.drawString(.82*inch, 9.55*inch, data["name"])
        canvas.setFillColor(HexColor("#DDD6FE"))
        canvas.setFont("Helvetica", 8.5)
        canvas.drawString(.82*inch, 9.2*inch, data["title"].upper())
        canvas.restoreState()
    document.addPageTemplates(PageTemplate(id="creative", frames=[frame], onPage=background))
    story = []
    if data["bio"]:
        story += [Paragraph("MON PARCOURS", styles["heading"]), Paragraph(_xml(data["bio"]), styles["body"])]
    if data["experiences"]:
        card_content = [Paragraph("EXPÉRIENCE RÉCENTE", styles["heading"])] + _pdf_experiences(data, styles, "#7C3AED", limit=3)
        card = Table([[card_content]], colWidths=[7.0*inch])
        card.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), HexColor("#F5F3FF")), ("BOX", (0,0), (-1,-1), 0, HexColor("#F5F3FF")), ("LEFTPADDING", (0,0), (-1,-1), 12), ("RIGHTPADDING", (0,0), (-1,-1), 12), ("TOPPADDING", (0,0), (-1,-1), 8), ("BOTTOMPADDING", (0,0), (-1,-1), 8)]))
        story += [card, Spacer(1, 7)]
    if data["skills"]:
        story.append(Paragraph("MES POINTS FORTS", styles["heading"]))
        cells = [Paragraph(_xml(skill), ParagraphStyle("Chip", parent=styles["small"], alignment=TA_CENTER, textColor=HexColor("#4C1D95"))) for skill in data["skills"][:9]]
        rows = [cells[index:index+3] for index in range(0, len(cells), 3)]
        while rows and len(rows[-1]) < 3:
            rows[-1].append("")
        chips = Table(rows, colWidths=[2.28*inch]*3, hAlign="LEFT")
        chips.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), HexColor("#DDD6FE")), ("GRID", (0,0), (-1,-1), 5, white), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5)]))
        story += [chips, Spacer(1, 4)]
    if data["education"]:
        story.append(Paragraph("FORMATION", styles["heading"]))
        story.extend(_pdf_education(data, styles, "#7C3AED"))
    document.build(story)
    buffer.seek(0)
    return buffer


def _pdf_cv_libre(generation):
    data = _cv_data(generation)
    styles = _pdf_styles("#4F46E5")
    header_style = ParagraphStyle("FreeName", parent=styles["name"], fontSize=20, textColor=HexColor("#1D4ED8"))
    header = Table([[[Paragraph(_xml(data["name"]), header_style), Paragraph(_xml(data["title"].upper()), styles["title"]), Paragraph(_xml(_doc_contact_line(data)), styles["small"])]]], colWidths=[7*inch])
    header.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), HexColor("#DBEAFE")), ("BOX", (0,0), (-1,-1), 0, white), ("LEFTPADDING", (0,0), (-1,-1), 14), ("RIGHTPADDING", (0,0), (-1,-1), 14), ("TOPPADDING", (0,0), (-1,-1), 10), ("BOTTOMPADDING", (0,0), (-1,-1), 10)]))
    story = [header, Spacer(1, 10)]
    profile = [Paragraph("PROFIL", styles["heading"]), Paragraph(_xml(data["bio"] or "Présentation professionnelle"), styles["small"])]
    skills = [Paragraph("COMPÉTENCES", styles["heading"]), Paragraph(_xml(" | ".join(data["skills"][:8]) or "Compétences du profil"), styles["small"])]
    overview = Table([[profile, skills]], colWidths=[3.45*inch, 3.45*inch])
    overview.setStyle(TableStyle([("BACKGROUND", (0,0), (0,0), HexColor("#E9D5FF")), ("BACKGROUND", (1,0), (1,0), HexColor("#BFDBFE")), ("GRID", (0,0), (-1,-1), 5, white), ("VALIGN", (0,0), (-1,-1), "TOP"), ("LEFTPADDING", (0,0), (-1,-1), 12), ("RIGHTPADDING", (0,0), (-1,-1), 12), ("TOPPADDING", (0,0), (-1,-1), 8), ("BOTTOMPADDING", (0,0), (-1,-1), 8)]))
    story += [overview, Spacer(1, 10)]
    if data["experiences"]:
        experience = [Paragraph("EXPÉRIENCE", styles["heading"])] + _pdf_experiences(data, styles, "#16A34A", limit=4)
        card = Table([[experience]], colWidths=[7*inch])
        card.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), HexColor("#DCFCE7")), ("BOX", (0,0), (-1,-1), 0, white), ("LEFTPADDING", (0,0), (-1,-1), 12), ("RIGHTPADDING", (0,0), (-1,-1), 12), ("TOPPADDING", (0,0), (-1,-1), 8), ("BOTTOMPADDING", (0,0), (-1,-1), 8)]))
        story += [card, Spacer(1, 8)]
    if data["education"]:
        story.append(Paragraph("FORMATION", styles["heading"]))
        story.extend(_pdf_education(data, styles, "#4F46E5"))
    buffer = BytesIO()
    pdf = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=.65*inch, leftMargin=.65*inch, topMargin=.55*inch, bottomMargin=.55*inch, title=generation.titre, author=data["name"])
    pdf.build(story)
    buffer.seek(0)
    return buffer


def _pdf_generic(generation):
    theme = _theme(generation.design)
    accent = "#" + theme["accent"]
    styles = _pdf_styles(accent, theme["pdf_font"])
    styles["name"].alignment = TA_CENTER if theme["title_align"] == "center" else TA_LEFT
    story = [Paragraph(_xml(generation.user.full_name or generation.titre), styles["name"])]
    if generation.user.email:
        story.append(Paragraph(_xml(generation.user.email), styles["title"]))
    for kind, text in _blocks(generation.contenu):
        if kind == "space":
            story.append(Spacer(1, 3))
        elif kind == "title" and _normalise(text) in _normalise(generation.user.full_name or generation.titre):
            continue
        elif kind in {"title", "heading", "subheading"}:
            story.append(Paragraph(_xml(text.upper() if kind == "heading" else text), styles["heading"]))
        elif kind == "bullet":
            story.append(Paragraph("- " + _xml(text), styles["body"]))
        else:
            story.append(Paragraph(_xml(text), styles["body"]))
    buffer = BytesIO()
    pdf = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=.7*inch, leftMargin=.7*inch, topMargin=.6*inch, bottomMargin=.6*inch, title=generation.titre, author=generation.user.full_name)
    pdf.build(story)
    buffer.seek(0)
    return buffer


def build_pdf(generation):
    builders = {
        "cv-classique": _pdf_cv_classique,
        "cv-moderne": _pdf_cv_moderne,
        "cv-creatif": _pdf_cv_creatif,
        "libre": _pdf_cv_libre if generation.outil == "generer-cv" else _pdf_generic,
    }
    return builders.get(generation.design, _pdf_generic)(generation)


def export_generation(generation, output_format):
    if output_format == "pdf":
        return build_pdf(generation), "application/pdf", ".pdf"
    if output_format == "docx":
        return build_docx(generation), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"
    raise ValueError("Format non pris en charge")
