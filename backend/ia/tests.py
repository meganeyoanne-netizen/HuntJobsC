import json
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.test import SimpleTestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

from .extraction import extract_text
from .services import complete


class CompleteTests(SimpleTestCase):
    @override_settings(GROQ_API_KEY="test-key", GROQ_MODEL="test-model")
    @patch("ia.services.urllib.request.urlopen")
    def test_request_uses_application_user_agent(self, urlopen):
        response = MagicMock()
        response.__enter__.return_value = response
        response.__exit__.return_value = False
        response.read.return_value = json.dumps(
            {"choices": [{"message": {"content": "CV généré"}}]}
        ).encode()
        urlopen.return_value = response

        result = complete("generer-cv", {"texte": "Profil"})

        request = urlopen.call_args.args[0]
        self.assertEqual(request.get_header("User-agent"), "HuntJobs/1.0")
        self.assertEqual(request.get_header("Accept"), "application/json")
        self.assertEqual(result, "CV généré")


class DocumentExtractionTests(SimpleTestCase):
    def test_extracts_docx_text(self):
        stream = BytesIO()
        document = Document()
        document.add_heading("Expérience", level=1)
        document.add_paragraph("Développement Django et React")
        document.save(stream)
        uploaded = SimpleUploadedFile(
            "cv.docx",
            stream.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        text = extract_text(uploaded)

        self.assertIn("Expérience", text)
        self.assertIn("Développement Django et React", text)

    @patch("legacy_doc.extract_text")
    def test_extracts_legacy_doc_text(self, legacy_extract):
        legacy_extract.return_value = SimpleNamespace(text="Ancien CV Word lisible")
        uploaded = SimpleUploadedFile("cv.doc", b"legacy-word-content", content_type="application/msword")

        text = extract_text(uploaded)

        legacy_extract.assert_called_once_with(b"legacy-word-content")
        self.assertEqual(text, "Ancien CV Word lisible")
